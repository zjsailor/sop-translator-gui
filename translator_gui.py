import io
import re
import threading
import time
import traceback
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import requests
from docx import Document
from openpyxl import load_workbook

APP_TITLE = "SOP保留格式翻译工具"
APP_VERSION = "Version 1.0"
APP_CONTACT = "hongxt"
APP_PHONE = "zjubell#hotmail.com"

BATCH_SIZE = 50
SEP = "|||---SEP---|||"


def contains_chinese(text: str) -> bool:
    return bool(re.search(r"[\u4e00-\u9fff]", text or ""))


def query_llm(base_url: str, api_key: str, model: str, prompt_text: str, timeout=120) -> str:
    payload = {
        "model": model,
        "messages": [
            {"role": "system", "content": "你是专业翻译助手。"},
            {"role": "user", "content": prompt_text},
        ],
        "temperature": 0.1,
    }
    headers = {"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"}
    resp = requests.post(base_url, json=payload, headers=headers, timeout=timeout)
    resp.raise_for_status()
    data = resp.json()
    return data["choices"][0]["message"]["content"].strip()


def translate_texts_serial(texts: list[str], base_url: str, api_key: str, model: str, target_lang: str) -> list[str]:
    out = []
    for t in texts:
        try:
            prompt = f"Translate to {target_lang}, only output translation:\n{t}"
            out.append(query_llm(base_url, api_key, model, prompt, timeout=60))
            time.sleep(0.1)
        except Exception:
            out.append(t)
    return out


def translate_texts_batch(texts: list[str], base_url: str, api_key: str, model: str, target_lang="英文") -> list[str]:
    if not texts:
        return []
    instruction = (
        f"请将下面中文逐条翻译为{target_lang}。\n"
        f"1. 严格保持顺序。\n"
        f"2. 仅输出翻译结果，不要额外解释。\n"
        f"3. 必须用分隔符 {SEP} 分隔每条结果。\n\n"
        "内容：\n"
    )
    prompt = instruction + SEP.join(texts)
    try:
        raw = query_llm(base_url, api_key, model, prompt)
        parts = [p.strip() for p in raw.split(SEP)]
        if len(parts) != len(texts):
            return translate_texts_serial(texts, base_url, api_key, model, target_lang)
        return parts
    except Exception:
        return translate_texts_serial(texts, base_url, api_key, model, target_lang)


def process_docx_object(doc_obj, text_list, target_list):
    for para in getattr(doc_obj, "paragraphs", []):
        txt = para.text.strip()
        if txt and contains_chinese(txt):
            text_list.append(txt)
            target_list.append(para)
    for table in getattr(doc_obj, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                process_docx_object(cell, text_list, target_list)


def extract_docx(doc):
    texts, targets = [], []
    process_docx_object(doc, texts, targets)
    for section in doc.sections:
        for h in [section.header, section.first_page_header, section.even_page_header]:
            if h and not h.is_linked_to_previous:
                process_docx_object(h, texts, targets)
        for f in [section.footer, section.first_page_footer, section.even_page_footer]:
            if f and not f.is_linked_to_previous:
                process_docx_object(f, texts, targets)
    return texts, targets


def apply_docx(targets, translations):
    for obj, trans in zip(targets, translations):
        if hasattr(obj, "runs") and len(obj.runs) > 0:
            obj.runs[0].text = trans
            for i in range(1, len(obj.runs)):
                obj.runs[i].text = ""
        else:
            obj.text = trans


def extract_excel(wb):
    texts, targets = [], []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and contains_chinese(cell.value):
                    texts.append(cell.value.strip())
                    targets.append(cell)
    return texts, targets


def process_single_file(input_file: Path, output_file: Path, base_url: str, api_key: str, model: str, log):
    ext = input_file.suffix.lower()
    if ext not in (".docx", ".xlsx"):
        raise ValueError("仅支持 .docx 或 .xlsx")

    log(f"正在处理: {input_file.name}")
    with open(input_file, "rb") as f:
        file_stream = io.BytesIO(f.read())

    if ext == ".docx":
        doc = Document(file_stream)
        texts, targets = extract_docx(doc)
    else:
        wb = load_workbook(file_stream)
        texts, targets = extract_excel(wb)

    translated = []
    for i in range(0, len(texts), BATCH_SIZE):
        log(f"翻译进度: {i}/{len(texts)}")
        translated.extend(translate_texts_batch(texts[i : i + BATCH_SIZE], base_url, api_key, model))

    if ext == ".docx":
        apply_docx(targets, translated)
        doc.save(output_file)
    else:
        for cell, trans in zip(targets, translated):
            cell.value = trans
        wb.save(output_file)

    log(f"输出完成: {output_file}")


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(APP_TITLE)
        self.geometry("850x620")

        self.base_url_var = tk.StringVar(value="https://api.deepseek.com/chat/completions")
        self.model_var = tk.StringVar(value="deepseek-chat")
        self.key_var = tk.StringVar()
        self.input_var = tk.StringVar()
        self.out_var = tk.StringVar(value=str((Path.home() / "Desktop").resolve()))
        self._build()

    def _build(self):
        frm = ttk.Frame(self, padding=16)
        frm.pack(fill="both", expand=True)

        top = ttk.Frame(frm)
        top.grid(row=0, column=0, columnspan=3, sticky="e")
        ttk.Button(top, text="About", command=self.show_about).pack(side="right")

        ttk.Label(frm, text="Base URL").grid(row=1, column=0, sticky="w", pady=(8, 0))
        ttk.Entry(frm, textvariable=self.base_url_var, width=95).grid(row=2, column=0, columnspan=3, sticky="we")

        ttk.Label(frm, text="AI 模型").grid(row=3, column=0, sticky="w", pady=(8, 0))
        ttk.Entry(frm, textvariable=self.model_var, width=35).grid(row=4, column=0, sticky="w")

        ttk.Label(frm, text="API Key (sk-...)").grid(row=3, column=1, sticky="w", pady=(8, 0))
        ttk.Entry(frm, textvariable=self.key_var, show="*", width=55).grid(row=4, column=1, columnspan=2, sticky="we")

        ttk.Label(frm, text="输入文件 (.docx/.xlsx)").grid(row=5, column=0, sticky="w", pady=(8, 0))
        ttk.Entry(frm, textvariable=self.input_var, width=80).grid(row=6, column=0, columnspan=2, sticky="we")
        ttk.Button(frm, text="选择", command=self.pick_input).grid(row=6, column=2, sticky="e")

        ttk.Label(frm, text="输出目录 (默认桌面)").grid(row=7, column=0, sticky="w", pady=(8, 0))
        ttk.Entry(frm, textvariable=self.out_var, width=80).grid(row=8, column=0, columnspan=2, sticky="we")
        ttk.Button(frm, text="选择", command=self.pick_out).grid(row=8, column=2, sticky="e")

        self.run_btn = ttk.Button(frm, text="转换", command=self.run_task)
        self.run_btn.grid(row=9, column=0, columnspan=3, sticky="we", pady=(10, 10))

        self.log_txt = tk.Text(frm, height=22, wrap="word")
        self.log_txt.grid(row=10, column=0, columnspan=3, sticky="nsew")

        frm.columnconfigure(0, weight=1)
        frm.columnconfigure(1, weight=1)
        frm.rowconfigure(10, weight=1)

    def show_about(self):
        messagebox.showinfo("About", f"{APP_VERSION}\n联系人: {APP_CONTACT}\n联系方式: {APP_PHONE}")

    def log(self, msg: str):
        self.log_txt.insert("end", f"{datetime.now().strftime('%H:%M:%S')}  {msg}\n")
        self.log_txt.see("end")
        self.update_idletasks()

    def pick_input(self):
        p = filedialog.askopenfilename(filetypes=[("Office", "*.docx *.xlsx")])
        if p:
            self.input_var.set(p)

    def pick_out(self):
        p = filedialog.askdirectory()
        if p:
            self.out_var.set(p)

    def run_task(self):
        base_url = self.base_url_var.get().strip()
        model = self.model_var.get().strip()
        api_key = self.key_var.get().strip()
        input_path = Path(self.input_var.get().strip())
        out_dir = Path(self.out_var.get().strip())

        if not base_url or not model or not api_key.startswith("sk-"):
            messagebox.showerror("错误", "请填写有效的 base_url、模型、sk-开头 API key。")
            return
        if not input_path.exists():
            messagebox.showerror("错误", "请选择有效输入文件。")
            return
        out_dir.mkdir(parents=True, exist_ok=True)

        output_file = out_dir / f"{input_path.stem}_EN{input_path.suffix}"
        self.run_btn.configure(state="disabled")
        self.log_txt.delete("1.0", "end")
        self.log("开始转换...")

        def worker():
            try:
                process_single_file(input_path, output_file, base_url, api_key, model, self.log)
                self.log("转换完成")
                messagebox.showinfo("完成", f"输出文件:\n{output_file}")
            except Exception as e:
                self.log("失败: " + str(e))
                self.log(traceback.format_exc())
                messagebox.showerror("失败", str(e))
            finally:
                self.run_btn.configure(state="normal")

        threading.Thread(target=worker, daemon=True).start()


if __name__ == "__main__":
    App().mainloop()
